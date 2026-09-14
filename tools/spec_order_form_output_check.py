import base64
import html
import io
import math
import re
import unicodedata
from collections import defaultdict

import fitz
import pandas as pd
import streamlit as st

from PIL import Image, ImageDraw, ImageFont, ImageChops

# ============================================================================
# TOOL 3 — ORG SPEC + ORDER FORM + OUTPUT QC
#
# Design rules:
#   1. Tool 1 remains the source of truth for Order Form -> Output checking.
#   2. Tool 3 adds ORG-first analysis, static/variable classification,
#      evidence locking, registration, visual comparison and reporting.
#   3. No second copy of Tool 1 field-validation logic is maintained here.
#   4. Variable evidence is classified/locked before static matching so that
#      one artwork occurrence cannot be counted twice.
#   5. Overlay/Blink use one registered output coordinate space.
# ============================================================================

try:
    from .order_form_output_check import (
        AUTO_DETECT_ENGINE_VERSION as TOOL1_ENGINE_VERSION,
        _apply_tool_css as _apply_tool1_css,
        build_report,
        check_field,
        extract_output_pages,
        get_available_fields,
        get_field_region,
        get_field_type,
        is_admin_field,
        is_blank_value,
        load_excel,
        normalize_text,
        auto_detect_fields,
    )
except Exception:
    from order_form_output_check import (
        AUTO_DETECT_ENGINE_VERSION as TOOL1_ENGINE_VERSION,
        _apply_tool_css as _apply_tool1_css,
        build_report,
        check_field,
        extract_output_pages,
        get_available_fields,
        get_field_region,
        get_field_type,
        is_admin_field,
        is_blank_value,
        load_excel,
        normalize_text,
        auto_detect_fields,
    )

TOOL3_VERSION = "2026-09-14-TOOL3-FIVE-STAGE-WIZARD-TOOL1-LINKED-V6"

# -------------------------------- palette ----------------------------------
BLUE = "#2563eb"
BLUE_DARK = "#0f4fbf"
CYAN = "#06b6d4"
PAGE_BG = "#f4f8fc"
CARD_BG = "#ffffff"
CARD_BORDER = "#d9e4f0"
TEXT = "#12233a"
MUTED = "#64748b"
GREEN = "#16a34a"
GREEN_BG = "#dcfce7"
RED = "#dc2626"
RED_BG = "#fee2e2"
ORANGE = "#ea580c"
ORANGE_BG = "#ffedd5"
PURPLE = "#9333ea"
PURPLE_BG = "#f3e8ff"
YELLOW = "#eab308"
YELLOW_BG = "#fef9c3"
GRAY_BG = "#eef2f7"

FIELD_COLORS = {
    "CARE": "#0891b2",
    "CONTENT": "#9333ea",
    "COO": "#0f766e",
    "RN": "#d97706",
    "CA": "#d97706",
    "IDENTIFIER": "#e11d48",
    "SIZE": "#4f46e5",
    "COLOR": "#db2777",
    "GENDER": "#0284c7",
    "BRAND": "#65a30d",
    "ATTRIBUTE": "#ca8a04",
    "QUANTITY": "#0f766e",
    "BATCH": "#ea580c",
    "BARCODE": "#0891b2",
    "OSZ": "#7c3aed",
    "SYMBOL": "#c026d3",
    "PRODUCTION_MARK": "#4f46e5",
    "GENERAL": "#64748b",
}


def _field_color(field):
    return FIELD_COLORS.get(get_field_type(field), FIELD_COLORS["GENERAL"])


def _status_badge(status):
    value = str(status or "INFO").upper()
    mapping = {
        "PASS": (GREEN_BG, GREEN),
        "FAIL": (RED_BG, RED),
        "REVIEW": (ORANGE_BG, ORANGE),
        "STATIC": (YELLOW_BG, "#92400e"),
        "VARIABLE": ("#dbeafe", BLUE_DARK),
        "MISSING / UNACCOUNTED": (PURPLE_BG, PURPLE),
        "LOCKED": ("#e0f2fe", "#0369a1"),
        "AUTO": ("#dcfce7", "#15803d"),
        "MANUAL": ("#e0e7ff", "#4338ca"),
        "INFO": (GRAY_BG, TEXT),
    }
    bg, fg = mapping.get(value, mapping["INFO"])
    return f'<span class="t3-badge" style="background:{bg};color:{fg}">{html.escape(value)}</span>'


# ============================================================================
# CSS / UI shell
# ============================================================================


def _tool3_css():
    # Do not allow Tool 1's dark dashboard CSS to control this screen. Tool 3
    # intentionally has its own clean workspace matching the approved mockup.
    st.markdown(
        f"""
        <style>
        html, body, [data-testid="stAppViewContainer"], [data-testid="stApp"],
        [data-testid="stMain"], .stApp, .main, .block-container {{
            background:{PAGE_BG} !important;
            color:{TEXT} !important;
        }}
        [data-testid="stHeader"] {{ background:{PAGE_BG} !important; }}
        .block-container {{ max-width:1560px !important; padding-top:0.8rem !important; padding-bottom:2rem !important; }}
        .stApp, .stApp p, .stApp label, .stApp span, .stApp div {{ color:{TEXT}; }}

        .t3-topbar {{
            display:flex; align-items:center; gap:14px; padding:6px 4px 14px 4px;
            border-bottom:1px solid {CARD_BORDER}; margin-bottom:16px;
        }}
        .t3-logo {{
            width:42px; height:42px; border-radius:12px; display:flex; align-items:center; justify-content:center;
            color:white !important; font-size:21px; font-weight:900;
            background:linear-gradient(135deg,#1d4ed8,#06b6d4); box-shadow:0 4px 15px rgba(37,99,235,.20);
        }}
        .t3-title {{ font-size:28px; font-weight:850; line-height:1.05; color:{TEXT} !important; }}
        .t3-subtitle {{ font-size:11px; color:{MUTED} !important; margin-top:4px; }}
        .t3-top-pills {{ margin-left:auto; display:flex; gap:7px; align-items:center; }}
        .t3-top-pill {{ background:#ffffff; border:1px solid {CARD_BORDER}; border-radius:999px; padding:7px 11px; font-size:10px; color:{MUTED} !important; }}

        .t3-stepper {{
            display:grid; grid-template-columns:repeat(5,1fr); gap:7px; margin:2px 0 14px 0;
        }}
        .t3-step {{ background:#fff; border:1px solid {CARD_BORDER}; border-radius:10px; padding:8px 10px; }}
        .t3-step.active {{ border-color:#93c5fd; background:#eff6ff; }}
        .t3-step.done {{ background:#f0fdf4; border-color:#bbf7d0; }}
        .t3-step-num {{ font-size:9px; font-weight:800; color:{BLUE}; }}
        .t3-step-name {{ font-size:11px; font-weight:800; margin-top:2px; }}
        .t3-step-note {{ font-size:8px; color:{MUTED} !important; margin-top:2px; }}

        .t3-card {{ background:{CARD_BG}; border:1px solid {CARD_BORDER}; border-radius:12px; padding:13px; box-shadow:0 1px 2px rgba(15,23,42,.03); }}
        .t3-card-title {{ font-size:13px; font-weight:820; color:{TEXT} !important; }}
        .t3-card-sub {{ font-size:9px; color:{MUTED} !important; line-height:1.45; }}
        .t3-upload-icon {{ font-size:22px; margin-bottom:5px; }}
        .t3-fileline {{ font-size:9px; color:#334155 !important; margin-top:7px; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }}
        .t3-role {{ margin-top:7px; background:#f8fbff; border:1px solid #e2ebf5; border-radius:8px; padding:7px; font-size:8px; color:{MUTED} !important; line-height:1.45; }}
        .t3-role b {{ color:{TEXT} !important; }}

        .t3-section-title {{ font-size:14px; font-weight:850; margin:13px 0 7px 0; color:{TEXT} !important; }}
        .t3-section-note {{ font-size:9px; color:{MUTED} !important; margin-bottom:7px; }}
        .t3-mini-stat {{ background:#fff; border:1px solid {CARD_BORDER}; border-radius:9px; padding:9px 10px; }}
        .t3-mini-stat b {{ display:block; font-size:19px; font-weight:850; }}
        .t3-mini-stat span {{ font-size:8px; color:{MUTED} !important; }}
        .t3-badge {{ display:inline-block; padding:3px 7px; border-radius:999px; font-size:8px; font-weight:850; white-space:nowrap; }}
        .t3-lock {{ color:#0369a1; font-size:11px; }}
        .t3-search-note {{ font-size:8px; color:{MUTED} !important; margin-top:4px; }}
        .t3-field-row {{ background:#fff; border:1px solid {CARD_BORDER}; border-radius:8px; padding:8px 9px; margin:4px 0; }}
        .t3-field-name {{ font-size:10px; font-weight:800; }}
        .t3-field-sub {{ font-size:8px; color:{MUTED} !important; margin-top:2px; }}
        .t3-finding {{ background:#fff; border:1px solid {CARD_BORDER}; border-radius:10px; padding:9px; margin-bottom:7px; }}
        .t3-finding.selected {{ border-color:#60a5fa; box-shadow:0 0 0 2px #dbeafe; }}
        .t3-finding-head {{ display:flex; justify-content:space-between; gap:8px; align-items:center; }}
        .t3-finding-field {{ font-size:10px; font-weight:850; }}
        .t3-finding-detail {{ font-size:8px; color:#64748b !important; line-height:1.5; margin-top:5px; }}
        .t3-lock-pill {{ font-size:8px; color:#0369a1 !important; background:#e0f2fe; border-radius:999px; padding:2px 6px; }}
        .t3-hero-result {{ background:#eff6ff; border:1px solid #bfdbfe; border-radius:12px; padding:12px; }}
        .t3-hero-title {{ font-size:11px; font-weight:850; }}
        .t3-hero-sub {{ font-size:8px; color:{MUTED} !important; margin-top:4px; line-height:1.45; }}
        .t3-bottom-note {{ text-align:center; color:#94a3b8 !important; font-size:9px; margin-top:12px; }}

        [data-baseweb="select"] > div {{ background:#fff !important; border:1px solid #bfd0e2 !important; border-radius:8px !important; min-height:36px !important; }}
        [data-baseweb="select"] input, [data-baseweb="select"] span {{ color:{TEXT} !important; }}
        [role="option"] {{ background:#fff !important; color:{TEXT} !important; }}
        [role="option"]:hover {{ background:#eff6ff !important; }}
        [data-baseweb="tag"] {{ background:#dbeafe !important; color:#1e40af !important; }}
        [data-baseweb="tag"] span {{ color:#1e40af !important; }}
        [data-testid="stTextInput"] input {{ background:#fff !important; color:{TEXT} !important; border:1px solid #bfd0e2 !important; border-radius:8px !important; }}
        [data-testid="stFileUploader"] {{ background:#fff !important; border:1px solid #cbd8e6 !important; border-radius:10px !important; padding:3px !important; }}
        [data-testid="stFileUploaderDropzone"] {{ background:#f8fbff !important; border:1px dashed #aac4df !important; border-radius:8px !important; }}
        [data-testid="stDataFrame"] {{ border:1px solid {CARD_BORDER} !important; border-radius:10px !important; }}
        div.stButton > button {{ background:#fff !important; color:#1e3a5f !important; border:1px solid #b9cce0 !important; border-radius:8px !important; font-weight:750 !important; min-height:36px !important; box-shadow:none !important; }}
        div.stButton > button:hover {{ border-color:#60a5fa !important; background:#eff6ff !important; }}
        div.stButton > button[kind="primary"] {{ background:linear-gradient(90deg,#2563eb,#4f46e5) !important; color:#fff !important; border:0 !important; }}
        div.stDownloadButton > button {{ background:#2563eb !important; color:#fff !important; border:0 !important; border-radius:8px !important; font-weight:800 !important; }}
        [data-testid="stExpander"] {{ background:#fff !important; border:1px solid {CARD_BORDER} !important; border-radius:9px !important; }}
        .stTabs [data-baseweb="tab-list"] {{ gap:5px; background:#eaf0f7; padding:4px; border-radius:8px; }}
        .stTabs [data-baseweb="tab"] {{ background:transparent; border-radius:6px; font-size:10px; }}
        .stTabs [aria-selected="true"] {{ background:#fff !important; box-shadow:0 1px 3px rgba(15,23,42,.08); }}
        </style>
        """,
        unsafe_allow_html=True,
    )


def _render_stepper(active_step=3):
    steps = [
        ("1", "Upload", "Files & Product Type"),
        ("2", "Detect", "Order Form fields"),
        ("3", "Review", "Lock evidence"),
        ("4", "Compare", "ORG + Output"),
        ("5", "Results", "QC Summary"),
    ]
    cells = []
    for idx, (num, name, note) in enumerate(steps, start=1):
        state = "done" if idx < active_step else ("active" if idx == active_step else "")
        cells.append(f'<div class="t3-step {state}"><div class="t3-step-num">{num}</div><div class="t3-step-name">{html.escape(name)}</div><div class="t3-step-note">{html.escape(note)}</div></div>')
    st.markdown('<div class="t3-stepper">' + ''.join(cells) + '</div>', unsafe_allow_html=True)


def _render_top():
    st.markdown(
        """
        <div class="t3-topbar">
          <div class="t3-logo">✓</div>
          <div>
            <div class="t3-title">Tool 3 — ORG + Order Form + Output QC</div>
            <div class="t3-subtitle">Smarter workflow • Unified engine from Tool 1 • Clear visual comparison • Locked evidence</div>
          </div>
          <div class="t3-top-pills">
            <div class="t3-top-pill">Tool 1 linked</div>
            <div class="t3-top-pill">Evidence locked</div>
            <div class="t3-top-pill">Registered viewer</div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================================
# Text / geometry helpers
# ============================================================================


def _norm(text):
    value = unicodedata.normalize("NFKC", str(text or "")).strip()
    value = re.sub(r"\s+", " ", value)
    return value.casefold()


def _compact(text):
    return re.sub(r"[^a-z0-9]+", "", _norm(text))


def _tokens(text):
    return re.findall(r"[a-z0-9%]+", _norm(text))


def _bbox_from_words(words):
    if not words:
        return None
    left = min(float(w.get("left", 0)) for w in words)
    top = min(float(w.get("top", 0)) for w in words)
    right = max(float(w.get("left", 0)) + float(w.get("width", 0)) for w in words)
    bottom = max(float(w.get("top", 0)) + float(w.get("height", 0)) for w in words)
    if right <= left or bottom <= top:
        return None
    return (left, top, right, bottom)


def _blocks(page):
    """Build visual blocks, preferring PDF direct words when available.

    Direct PDF word coordinates are substantially safer than OCR for ORG/static
    evidence because they preserve the document's own text and coordinates.
    Image-only artwork falls back to the OCR block stream.
    """
    direct = [w for w in (page or {}).get("direct_words", []) if str(w.get("text", "")).strip()]
    if direct:
        items = sorted(direct, key=lambda x: (float(x.get("top", 0)), float(x.get("left", 0))))
        lines = []
        for word in items:
            cx = float(word.get("left", 0))
            cy = float(word.get("top", 0)) + float(word.get("height", 0)) / 2
            h = max(1.0, float(word.get("height", 0)))
            placed = None
            for line in reversed(lines[-4:]):
                if abs(cy - line["cy"]) <= max(8.0, h * 0.48):
                    placed = line
                    break
            if placed is None:
                placed = {"cy": cy, "words": []}
                lines.append(placed)
            placed["words"].append(word)
        result = []
        for line in lines:
            ws = sorted(line["words"], key=lambda x: float(x.get("left", 0)))
            bbox = _bbox_from_words(ws)
            text = " ".join(str(x.get("text", "")).strip() for x in ws).strip()
            if not bbox or not re.search(r"[A-Za-z0-9]", text):
                continue
            l,t,r,b = bbox
            result.append({"text":text,"norm":_norm(text),"compact":_compact(text),"bbox":bbox,"cx":(l+r)/2,"cy":(t+b)/2,"width":r-l,"height":b-t,"words":ws})
        result.sort(key=lambda b:(b["bbox"][1],b["bbox"][0]))
        for i,b in enumerate(result): b["index"]=i
        return result

    words = [w for w in (page or {}).get("ocr_words", []) if str(w.get("text", "")).strip()]
    grouped = defaultdict(list)
    for word in words:
        key = (word.get("block_num", 0), word.get("par_num", 0), word.get("line_num", 0))
        grouped[key].append(word)
    blocks = []
    for key, items in grouped.items():
        items = sorted(items, key=lambda x: (float(x.get("top", 0)), float(x.get("left", 0))))
        bbox = _bbox_from_words(items)
        if not bbox:
            continue
        text = " ".join(str(x.get("text", "")).strip() for x in items).strip()
        if not text or not re.search(r"[A-Za-z0-9]", text):
            continue
        l, t, r, b = bbox
        blocks.append({"text": text, "norm": _norm(text), "compact": _compact(text), "bbox": bbox, "cx": (l+r)/2, "cy": (t+b)/2, "width": r-l, "height": b-t, "words": items})
    blocks.sort(key=lambda b: (b["bbox"][1], b["bbox"][0]))
    for i,b in enumerate(blocks): b["index"]=i
    return blocks


def _direct_text_boxes(page, target):
    """Return exact/compact text evidence from PDF direct-word coordinates."""
    direct = [w for w in (page or {}).get("direct_words", []) if str(w.get("text", "")).strip()]
    if not direct or not str(target or "").strip():
        return []
    target_n = _norm(target)
    target_c = _compact(target)
    # Group words into visual lines first.
    lines=[]
    for word in sorted(direct, key=lambda x:(float(x.get("top",0)),float(x.get("left",0)))):
        cy=float(word.get("top",0))+float(word.get("height",0))/2
        placed=None
        for line in reversed(lines[-4:]):
            h=max(1.0,float(word.get("height",0)))
            if abs(cy-line["cy"])<=max(8.0,h*.5): placed=line; break
        if placed is None:
            placed={"cy":cy,"words":[]}; lines.append(placed)
        placed["words"].append(word)
    found=[]
    for line in lines:
        ws=sorted(line["words"],key=lambda x:float(x.get("left",0)))
        text=" ".join(str(x.get("text","")).strip() for x in ws).strip()
        n=_norm(text); c=_compact(text)
        if n==target_n or c==target_c or (target_c and target_c in c):
            b=_bbox_from_words(ws)
            if b: found.append(b)
    return found


def _content_bbox(image):
    if image is None:
        return None
    rgb = image.convert("RGB")
    bg = Image.new("RGB", rgb.size, (255, 255, 255))
    diff = ImageChops.difference(rgb, bg)
    mask = diff.convert("L").point(lambda p: 0 if p >= 10 else 255)
    bbox = mask.getbbox()
    if not bbox:
        return (0, 0, rgb.width, rgb.height)
    l, t, r, b = bbox
    if r - l < rgb.width * .05 or b - t < rgb.height * .05:
        return (0, 0, rgb.width, rgb.height)
    return bbox


def _register(org_page, output_page):
    if not org_page or not output_page:
        return None
    org_raw = org_page.get("image_bytes")
    out_raw = output_page.get("image_bytes")
    if not org_raw or not out_raw:
        return None
    org_img = Image.open(io.BytesIO(org_raw)).convert("RGB")
    out_img = Image.open(io.BytesIO(out_raw)).convert("RGB")
    ob = _content_bbox(org_img)
    ub = _content_bbox(out_img)
    ow = max(1, ob[2] - ob[0]); oh = max(1, ob[3] - ob[1])
    uw = max(1, ub[2] - ub[0]); uh = max(1, ub[3] - ub[1])
    sx = uw / ow; sy = uh / oh
    # Uniform scaling is intentional: never stretch ORG in one axis.
    scale = max(.2, min(5.0, (sx + sy) / 2))
    org_c = ((ob[0] + ob[2]) / 2, (ob[1] + ob[3]) / 2)
    out_c = ((ub[0] + ub[2]) / 2, (ub[1] + ub[3]) / 2)
    tx = out_c[0] - org_c[0] * scale
    ty = out_c[1] - org_c[1] * scale
    aspect_delta = abs(sx - sy) / max(1e-6, (sx + sy) / 2)
    return {
        "org_image": org_img,
        "output_image": out_img,
        "org_bbox": ob,
        "output_bbox": ub,
        "scale": scale,
        "scale_x": sx,
        "scale_y": sy,
        "tx": tx,
        "ty": ty,
        "aspect_delta": aspect_delta,
        "warning": "Aspect ratio differs; uniform registration is used." if aspect_delta > .08 else "",
    }


def _transform_bbox(b, reg):
    l, t, r, bot = b
    s = reg["scale"]
    return (l * s + reg["tx"], t * s + reg["ty"], r * s + reg["tx"], bot * s + reg["ty"])


def _registered_org_blocks(org_page, reg):
    result = []
    for b in _blocks(org_page):
        nb = dict(b)
        nb["bbox"] = _transform_bbox(b["bbox"], reg)
        l, t, r, bot = nb["bbox"]
        nb["cx"] = (l + r) / 2
        nb["cy"] = (t + bot) / 2
        nb["width"] = r - l
        nb["height"] = bot - t
        result.append(nb)
    return result


def _iou(a, b):
    ax1, ay1, ax2, ay2 = a; bx1, by1, bx2, by2 = b
    ix1, iy1 = max(ax1, bx1), max(ay1, by1)
    ix2, iy2 = min(ax2, bx2), min(ay2, by2)
    iw, ih = max(0, ix2 - ix1), max(0, iy2 - iy1)
    inter = iw * ih
    if inter <= 0:
        return 0.0
    aa = max(1.0, (ax2-ax1)*(ay2-ay1)); bb = max(1.0, (bx2-bx1)*(by2-by1))
    return inter / (aa + bb - inter)


def _intersects_locked(box, locks, threshold=.16):
    return any(_iou(box, b) >= threshold for b in locks)


def _block_near_box(blocks, box, excluded=None):
    excluded = excluded or set()
    if not box:
        return None, None
    cx = (box[0] + box[2]) / 2; cy = (box[1] + box[3]) / 2
    bw = max(1.0, box[2]-box[0]); bh = max(1.0, box[3]-box[1])
    best = None; best_score = float("inf")
    for idx, b in enumerate(blocks):
        if idx in excluded:
            continue
        dx = abs(b["cx"]-cx) / max(1, bw*2)
        dy = abs(b["cy"]-cy) / max(1, bh*3)
        size = abs(math.log(max(.05, b["width"]) / max(.05, bw)))
        score = dx + dy + .15*size
        if score < best_score:
            best_score = score; best = (idx, b)
    return best if best else (None, None)


def _exact_static_score(org_block, out_block, page_w, page_h):
    if org_block["norm"] != out_block["norm"] and org_block["compact"] != out_block["compact"]:
        return None
    ox, oy = org_block["cx"], org_block["cy"]
    ux, uy = out_block["cx"], out_block["cy"]
    center = math.hypot((ox-ux)/max(1,page_w), (oy-uy)/max(1,page_h))
    if center > .05:
        return None
    wr = out_block["width"] / max(1, org_block["width"])
    hr = out_block["height"] / max(1, org_block["height"])
    if not (.50 <= wr <= 2.0 and .50 <= hr <= 2.0):
        return None
    score = (1-min(1, center/.05))*.65 + min(1, min(wr,1/wr))*0.2 + min(1,min(hr,1/hr))*.15
    return score


def _match_static(org_page, output_page, reg, locked_output_boxes, locked_org_indices):
    if not reg:
        return []
    org_blocks = _registered_org_blocks(org_page, reg)
    out_blocks = _blocks(output_page)
    page_w = max(1, int(output_page.get("image_width", 1)))
    page_h = max(1, int(output_page.get("image_height", 1)))
    candidates = []
    for oi, ob in enumerate(org_blocks):
        if oi in locked_org_indices or len(ob["compact"]) < 2:
            continue
        # Static means this exact ORG occurrence still has no variable evidence claim.
        for ui, ub in enumerate(out_blocks):
            if _intersects_locked(ub["bbox"], locked_output_boxes):
                continue
            score = _exact_static_score(ob, ub, page_w, page_h)
            if score is not None:
                candidates.append((score, oi, ui))
    candidates.sort(reverse=True)
    used_o = set(locked_org_indices); used_u = set()
    matches = []
    for score, oi, ui in candidates:
        if oi in used_o or ui in used_u:
            continue
        used_o.add(oi); used_u.add(ui)
        matches.append({"org": org_blocks[oi], "output": out_blocks[ui], "org_index": oi, "output_index": ui, "score": score})
    return matches


# ============================================================================
# Tool 1-linked variable evidence
# ============================================================================


def _page_row_mapping(df, output_pages, mapping=None):
    if mapping:
        return mapping
    if len(df) == 1:
        return {int(p.get("page", i+1)): 0 for i, p in enumerate(output_pages)}
    return {int(p.get("page", i+1)): min(i, len(df)-1) for i, p in enumerate(output_pages)}


def _tool1_variable_results(df, output_pages, selected_fields, product_type, mapping):
    """Run the real Tool 1 report. No duplicate field-validation logic."""
    if not selected_fields:
        return pd.DataFrame()
    # The actual comparison, field priority, specialized matchers, and report
    # semantics all come directly from Tool 1.
    return build_report(df, output_pages, selected_fields, product_type, page_row_mapping=mapping) if "page_row_mapping" in build_report.__code__.co_varnames else build_report(df, output_pages, selected_fields, product_type)


def _tool1_auto_detect(df, output_pages, product_type, mapping):
    try:
        return auto_detect_fields(df, output_pages, product_type, page_row_mapping=mapping)
    except TypeError:
        return auto_detect_fields(df, output_pages, product_type)


def _result_rows_by_page(report):
    if report is None or report.empty:
        return {}
    result = defaultdict(list)
    for _, row in report.iterrows():
        try:
            page = int(row.get("PDF PAGE", 1))
        except Exception:
            page = 1
        result[page].append(row.to_dict())
    return result


def _ocr_exact_scalar_boxes(page, target):
    """OCR fallback for short scalar values when Tool 1 deliberately avoids tiny token highlights."""
    target=_norm(target)
    compact=_compact(target)
    if not target or len(compact)<2:
        return []
    found=[]
    for w in (page or {}).get("ocr_words",[]) or []:
        text=str(w.get("text","")).strip()
        if not text:
            continue
        wn=_norm(text); wc=_compact(text)
        if wn==target or wc==compact:
            l=float(w.get("left",0)); t=float(w.get("top",0)); r=l+float(w.get("width",0)); b=t+float(w.get("height",0))
            if r>l and b>t: found.append((l,t,r,b))
    return found


def _ocr_contains_scalar_boxes(page, target):
    target=_compact(target)
    if not target or len(target)<2:
        return []
    found=[]
    for w in (page or {}).get("ocr_words",[]) or []:
        text=str(w.get("text","")).strip()
        c=_compact(text)
        if target and target in c:
            l=float(w.get("left",0)); t=float(w.get("top",0)); r=l+float(w.get("width",0)); b=t+float(w.get("height",0))
            if r>l and b>t: found.append((l,t,r,b))
    return found


def _ocr_subtoken_boxes(page, target):
    """Approximate a safe sub-box when OCR combines tokens such as 44-12."""
    target=str(target or "").strip()
    compact_target=_compact(target)
    if len(compact_target)<2: return []
    found=[]
    for w in (page or {}).get("ocr_words",[]) or []:
        text=str(w.get("text","")).strip()
        n=_norm(text); c=_compact(text)
        if not text or compact_target not in c: continue
        # Work in normalized displayed token string. If punctuation differs, use
        # the compact index and map it back to the original character positions.
        low=text.casefold(); c_low=_compact(text)
        ci=c_low.find(compact_target)
        if ci<0: continue
        compact_positions=[]
        for pos,ch in enumerate(low):
            if ch.isalnum(): compact_positions.append(pos)
        if ci >= len(compact_positions): continue
        start_pos=compact_positions[ci]
        end_idx=min(ci+len(compact_target)-1,len(compact_positions)-1)
        end_pos=compact_positions[end_idx]+1
        l=float(w.get("left",0)); t=float(w.get("top",0)); r=l+float(w.get("width",0)); b=t+float(w.get("height",0))
        char_count=max(1,len(text))
        # Width proportional estimate; generous padding keeps the token visible.
        sl=l + (start_pos/char_count)*(r-l)
        sr=l + (end_pos/char_count)*(r-l)
        pad=max(2,(r-l)*0.018)
        if sr>sl: found.append((max(l,sl-pad),t,min(r,sr+pad),b))
    return found


def _visual_boxes_for_result(output_page, row):
    field = str(row.get("FIELD", ""))
    expected = str(row.get("ORDER FORM DATA", "") or "")
    actual = str(row.get("PDF OUTPUT", "") or "")
    status = str(row.get("STATUS", ""))
    if status not in {"PASS", "FAIL"}:
        return []

    # Prefer the PDF text layer whenever one exists. This avoids OCR errors such
    # as F1607 -> 1FG1 on vector PDF artwork.
    direct=[]
    for target in (actual, expected):
        if target and target.casefold() not in {"not found","-","—"}:
            direct=_direct_text_boxes(output_page,target)
            if direct: break

    field_type=get_field_type(field)
    if direct and field_type not in {"CONTENT","CARE","SIZE","OSZ"}:
        return direct[:1]
    if direct and field_type in {"CONTENT","CARE","SIZE","OSZ"}:
        return direct

    # Short numeric/alphanumeric scalars get an exact OCR-word fallback. This is
    # intentionally blocked for one-character values to avoid the old isolated-S
    # / 0 / 1 false-highlighting problem.
    scalar_target = actual if actual and actual.casefold() not in {"not found","-","—"} else expected
    if field_type not in {"CONTENT","CARE","OSZ"}:
        scalar_boxes = _ocr_exact_scalar_boxes(output_page, scalar_target)
        if scalar_boxes:
            return scalar_boxes[:1]

    # Numeric SIZE values often live inside a combined line such as "44 - 12".
    # Prefer the exact OCR word first so Size and Size-Modifier can highlight
    # their own token rather than locking the entire combined line together.
    if field_type == "SIZE":
        scalar_boxes = _ocr_exact_scalar_boxes(output_page, scalar_target)
        if scalar_boxes:
            return scalar_boxes[:1]

    # Numeric SIZE values often live inside a combined line such as "44 - 12".
    # A safe contains fallback catches that exact token without enabling one-character
    # substring matching.
    if field_type == "SIZE":
        scalar_subtoken = _ocr_subtoken_boxes(output_page, scalar_target)
        if scalar_subtoken:
            return scalar_subtoken[:1]
        scalar_contains = _ocr_contains_scalar_boxes(output_page, scalar_target)
        if scalar_contains:
            return scalar_contains[:1]

    # Final semantic fallback: search the Tool 1-reported actual value using the
    # current page's PDF text/OCR coordinates. We intentionally do not maintain
    # a second field-matching engine here.
    semantic_targets = []
    if actual and actual.casefold() not in {"not found", "-", "—"}:
        semantic_targets.append(actual)
    if expected and expected.casefold() not in {"not found", "-", "—"}:
        semantic_targets.append(expected)

    boxes = []
    for target in semantic_targets:
        boxes = _direct_text_boxes(output_page, target)
        if boxes:
            break
        boxes = _ocr_contains_scalar_boxes(output_page, target)
        if boxes:
            break
    cleaned=[]
    for b in boxes or []:
        try:
            l,t,r,bot=[float(x) for x in b]
            if r>l and bot>t: cleaned.append((l,t,r,bot))
        except Exception:
            continue
    return cleaned


def _merge_evidence_boxes(boxes, gap=10):
    cleaned=[]
    for b in boxes or []:
        l,t,r,bot=[float(x) for x in b]
        if r>l and bot>t: cleaned.append((l,t,r,bot))
    cleaned.sort(key=lambda b:(b[1],b[0]))
    out=[]
    for b in cleaned:
        if not out:
            out.append(b); continue
        a=out[-1]
        vertical=min(a[3],b[3])-max(a[1],b[1])
        horizontal=max(0,max(a[0],b[0])-min(a[2],b[2]))
        if vertical>-gap and horizontal<=gap:
            out[-1]=(min(a[0],b[0]),min(a[1],b[1]),max(a[2],b[2]),max(a[3],b[3]))
        else:
            out.append(b)
    return out


def _visible_text_exact(reference, actual):
    """Exact visible-text equality for STATIC classification.

    Unlike Tool 1's case-insensitive validation, ORG STATIC classification is
    intentionally strict: case and punctuation are part of the visible baseline.
    """
    def canon(value):
        value = unicodedata.normalize("NFKC", str(value or "")).strip()
        value = re.sub(r"\s+", " ", value)
        return value
    return bool(canon(reference)) and canon(reference) == canon(actual)


def _build_variable_evidence(df, output_pages, org_pages, report, product_type):
    """Build one visual evidence object per Tool 1 result.

    Decision order:
      1. Tool 1 decides whether the Order Form field is PASS/FAIL.
      2. If the same visible Output occurrence exactly matches the registered ORG
         occurrence in case + punctuation + wording + position, it is STATIC.
      3. Otherwise it is VARIABLE and Tool 3 applies ORG presentation rules.
      4. Physical Output/ORG evidence is locked so later fields cannot reuse it.
    """
    rows_by_page = _result_rows_by_page(report)
    all_evidence = []

    for idx, out_page in enumerate(output_pages):
        page_no = int(out_page.get("page", idx + 1))
        page_rows = rows_by_page.get(page_no, [])
        org_page = org_pages[idx] if idx < len(org_pages) else None
        reg = _register(org_page, out_page) if org_page else None
        reg_org_blocks = _registered_org_blocks(org_page, reg) if reg else []
        output_blocks = _blocks(out_page)

        locked_out = []
        locked_org = set()

        for seq, row in enumerate(page_rows, start=1):
            field = str(row.get("FIELD", "")).strip()
            base_status = str(row.get("STATUS", "")).strip()
            if not field or base_status not in {"PASS", "FAIL"}:
                continue

            expected = str(row.get("ORDER FORM DATA", "") or "")
            actual = str(row.get("PDF OUTPUT", "") or "")
            boxes = _merge_evidence_boxes(_visual_boxes_for_result(out_page, row), gap=10)

            combined_box = None
            if boxes:
                combined_box = (
                    min(b[0] for b in boxes), min(b[1] for b in boxes),
                    max(b[2] for b in boxes), max(b[3] for b in boxes)
                )

            # Find the registered ORG region nearest to the Tool 1 evidence.
            org_idx = None
            org_block = None
            if combined_box and reg_org_blocks:
                org_idx, org_block = _block_near_box(reg_org_blocks, combined_box, locked_org)

            # Find the exact Output block corresponding to the evidence box.
            output_block = None
            if combined_box and output_blocks:
                out_idx, output_block = _block_near_box(output_blocks, combined_box)
                if output_block is None:
                    output_block = {
                        "text": actual or expected,
                        "bbox": combined_box,
                        "cx": (combined_box[0] + combined_box[2]) / 2,
                        "cy": (combined_box[1] + combined_box[3]) / 2,
                    }

            # ---------------------------------------------------------------
            # STATIC FIRST
            # ---------------------------------------------------------------
            static_exact = False
            static_score = None
            if org_block and output_block and _visible_text_exact(org_block.get("text", ""), output_block.get("text", "")):
                try:
                    static_score = _exact_static_score(
                        org_block,
                        output_block,
                        max(1, int(out_page.get("image_width", 1))),
                        max(1, int(out_page.get("image_height", 1))),
                    )
                except Exception:
                    static_score = None
                static_exact = static_score is not None

            if static_exact:
                status = "STATIC"
                classification = "STATIC"
                presentation_status = "PASS"
                presentation_reason = "Output exactly follows the ORG wording, case and punctuation at the registered location."
                difference = "Static ORG match."
                shared_evidence = False
                evidence_boxes = boxes
                locked = bool(boxes) or org_idx is not None

                if evidence_boxes:
                    for b in evidence_boxes:
                        locked_out.append(b)
                if org_idx is not None:
                    locked_org.add(org_idx)

                all_evidence.append({
                    "page": page_no,
                    "field": field,
                    "field_type": get_field_type(field),
                    "region": get_field_region(field),
                    "expected": expected,
                    "actual": actual,
                    "base_status": base_status,
                    "status": status,
                    "classification": classification,
                    "difference": difference,
                    "match_type": row.get("MATCH TYPE", row.get("match_type", "")),
                    "boxes": evidence_boxes,
                    "shared_evidence": shared_evidence,
                    "org_index": org_idx,
                    "org_block": org_block,
                    "output_block": output_block,
                    "locked": locked,
                    "source": "Tool 1 + ORG exact",
                    "sequence": seq,
                    "presentation_status": presentation_status,
                    "presentation_reason": presentation_reason,
                })
                continue

            # ---------------------------------------------------------------
            # VARIABLE / SHARED EVIDENCE
            # ---------------------------------------------------------------
            overlap_alias = bool(
                boxes and locked_out and
                all(_intersects_locked(box, locked_out, threshold=.45) for box in boxes)
            )
            evidence_boxes = [] if overlap_alias else boxes
            if evidence_boxes:
                for b in evidence_boxes:
                    locked_out.append(b)

            if overlap_alias:
                # If a second field points at the same physical occurrence, it is
                # linked to the existing evidence instead of drawing a second tag.
                shared_owner = None
                for previous in reversed(all_evidence):
                    if previous.get("page") != page_no:
                        continue
                    for pb in previous.get("boxes", []):
                        if any(_iou(pb, nb) >= .45 for nb in boxes):
                            shared_owner = previous.get("field")
                            break
                    if shared_owner:
                        break
                difference_prefix = f"Shared evidence with {shared_owner}. " if shared_owner else "Shared physical evidence. "
            else:
                difference_prefix = ""

            combined_for_org = None
            if evidence_boxes:
                combined_for_org = (
                    min(b[0] for b in evidence_boxes), min(b[1] for b in evidence_boxes),
                    max(b[2] for b in evidence_boxes), max(b[3] for b in evidence_boxes)
                )
            elif combined_box:
                combined_for_org = combined_box

            # Find the ORG block without consuming already-locked ORG evidence.
            org_idx = None
            org_block = None
            if combined_for_org and reg_org_blocks:
                org_idx, org_block = _block_near_box(reg_org_blocks, combined_for_org, locked_org)

            presentation_status, presentation_reason = _presentation_for_variable(
                org_block, actual, product_type
            )
            final_status = base_status
            difference = str(row.get("DIFFERENCE", "—") or "—")
            if presentation_status == "FAIL" and base_status == "PASS" and product_type in {"HTL", "Other"}:
                final_status = "FAIL"
                difference = (difference + " " + presentation_reason).strip()
            elif presentation_status == "REVIEW" and base_status == "PASS" and product_type in {"HTL", "Other"}:
                final_status = "REVIEW"
                difference = presentation_reason

            if difference_prefix:
                difference = difference_prefix + difference

            locked = bool(evidence_boxes) or org_idx is not None
            if org_idx is not None:
                locked_org.add(org_idx)

            all_evidence.append({
                "page": page_no,
                "field": field,
                "field_type": get_field_type(field),
                "region": get_field_region(field),
                "expected": expected,
                "actual": actual,
                "base_status": base_status,
                "status": final_status,
                "classification": "VARIABLE",
                "difference": difference,
                "match_type": row.get("MATCH TYPE", row.get("match_type", "")),
                "boxes": evidence_boxes,
                "shared_evidence": overlap_alias,
                "shared_owner": shared_owner if overlap_alias else None,
                "org_index": org_idx,
                "org_block": org_block,
                "output_block": output_block,
                "locked": locked,
                "source": "Tool 1",
                "sequence": seq,
                "presentation_status": presentation_status,
                "presentation_reason": presentation_reason,
            })

    return all_evidence


# ============================================================================
# ORG presentation / variable mapping
# ============================================================================


def _case_mode(word):
    letters = re.sub(r"[^A-Za-z]", "", str(word or ""))
    if not letters: return "NONE"
    if letters.isupper(): return "UPPER"
    if letters.islower(): return "LOWER"
    if len(letters)>1 and letters[0].isupper() and letters[1:].islower(): return "TITLE"
    return "MIXED"


def _presentation_for_variable(org_block, actual, product_type):
    if not org_block:
        return "REVIEW", "No registered ORG reference region was safely mapped to this variable field."
    org_text = org_block.get("text", "")
    actual = str(actual or "")
    if not org_text or not actual:
        return "REVIEW", "Insufficient presentation evidence."
    if product_type == "PFL":
        return "PASS", "PFL uses Order Form presentation for variable data."
    ref_words = re.findall(r"[A-Za-z]+", org_text)
    out_words = re.findall(r"[A-Za-z]+", actual)
    if ref_words and out_words:
        modes = [_case_mode(w) for w in ref_words]
        dominant = max(set(modes), key=modes.count)
        ratio = modes.count(dominant) / len(modes)
        if ratio >= .75:
            bad = [w for w in out_words if _case_mode(w) != dominant]
            if bad:
                return "FAIL", f"ORG establishes {dominant.lower()} presentation; Output does not preserve that case pattern."
    org_punct = [c for c in org_text if not c.isalnum() and not c.isspace()]
    out_punct = [c for c in actual if not c.isalnum() and not c.isspace()]
    if org_punct != out_punct:
        return "FAIL", f"ORG punctuation {org_punct or ['(none)']} differs from Output {out_punct or ['(none)']}."
    return "PASS", "ORG presentation requirements are consistent with Output."


# ============================================================================
# Static / unaccounted classification
# ============================================================================


def _classify_static_and_unaccounted(org_pages, output_pages, variable_evidence):
    """Classify remaining ORG blocks after variable/static field evidence is locked."""
    locked_output = defaultdict(list)
    locked_org = defaultdict(set)
    static_matches = []
    registrations = []
    unaccounted = []

    # First lock evidence that was itself proven to be an exact ORG STATIC field.
    for ev in variable_evidence:
        if ev.get("classification") != "STATIC":
            continue
        page_no = int(ev.get("page", 1))
        for b in ev.get("boxes", []):
            locked_output[page_no].append(b)
        if ev.get("org_index") is not None:
            locked_org[page_no].add(ev["org_index"])
        if ev.get("org_block") and ev.get("output_block"):
            static_matches.append({
                "page": page_no,
                "org": ev["org_block"],
                "output": ev["output_block"],
                "org_index": ev.get("org_index"),
                "output_index": ev.get("output_block", {}).get("index"),
                "status": "STATIC",
                "classification": "STATIC",
                "score": 1.0,
                "field": ev.get("field"),
            })

    for idx, out_page in enumerate(output_pages):
        page_no = int(out_page.get("page", idx + 1))
        org_page = org_pages[idx] if idx < len(org_pages) else None
        if not org_page:
            continue
        reg = _register(org_page, out_page)
        if not reg:
            continue
        registrations.append((page_no, reg))

        matches = _match_static(
            org_page,
            out_page,
            reg,
            locked_output.get(page_no, []),
            locked_org.get(page_no, set()),
        )
        for m in matches:
            # Avoid duplicate ORG index records if the same exact static evidence
            # was already claimed by a variable field classified as STATIC.
            if m.get("org_index") in locked_org.get(page_no, set()):
                continue
            static_matches.append({
                "page": page_no,
                "org": m["org"],
                "output": m["output"],
                "org_index": m.get("org_index"),
                "output_index": m.get("output_index"),
                "status": "STATIC",
                "classification": "STATIC",
                "score": m["score"],
            })
            locked_output[page_no].append(m["output"]["bbox"])
            locked_org[page_no].add(m["org_index"])

        org_blocks = _registered_org_blocks(org_page, reg)
        out_blocks = _blocks(out_page)
        static_org_idx = set(locked_org.get(page_no, set()))
        for oi, ob in enumerate(org_blocks):
            if oi in static_org_idx or len(ob.get("compact", "")) < 2:
                continue

            occupied = False
            for ub in out_blocks:
                if _intersects_locked(ub["bbox"], locked_output.get(page_no, []), threshold=.10):
                    occupied = True
                    break
            if occupied:
                continue

            obb = ob["bbox"]
            if (
                reg["output_bbox"][0] <= obb[0] <= reg["output_bbox"][2]
                and reg["output_bbox"][1] <= obb[1] <= reg["output_bbox"][3]
            ):
                unaccounted.append({
                    "page": page_no,
                    "field": "ORG Element",
                    "expected": ob["text"],
                    "actual": "Not found",
                    "status": "MISSING / UNACCOUNTED",
                    "difference": "ORG element has no safely mapped Output counterpart after variable/static evidence locking.",
                    "org": ob,
                    "output": None,
                })

    return static_matches, unaccounted, registrations


# ============================================================================
# Visual annotations / overlay viewer
# ============================================================================


def _rgba(hex_color, alpha):
    value = hex_color.lstrip("#")
    return tuple(int(value[i:i+2],16) for i in (0,2,4)) + (alpha,)


def _draw_visual_evidence(page, variable_evidence, static_matches, selected_field=None):
    raw = page.get("image_bytes") if page else None
    if not raw:
        return None
    base = Image.open(io.BytesIO(raw)).convert("RGBA")
    overlay = Image.new("RGBA", base.size, (0,0,0,0))
    draw = ImageDraw.Draw(overlay)
    entries = []

    page_no = int(page.get("page",1))
    # Static first, but only where variable evidence did not already claim the region.
    for idx, m in enumerate([x for x in static_matches if x["page"] == page_no], start=1):
        b = m["output"]["bbox"]
        draw.rounded_rectangle(tuple(map(int,b)), radius=4, fill=_rgba(YELLOW,70), outline=_rgba(YELLOW,235), width=2)
        entries.append((b, f"S{idx}", "STATIC", YELLOW))

    vars_page = [x for x in variable_evidence if x["page"] == page_no and x.get("classification", "VARIABLE") == "VARIABLE" and x.get("status") in {"PASS","FAIL","REVIEW"}]
    number_map = {x["field"]: i+1 for i,x in enumerate(variable_evidence)}
    occupied = []
    for ev in vars_page:
        for b in ev.get("boxes",[]):
            color = RED if ev["status"] == "FAIL" else ORANGE if ev["status"] == "REVIEW" else _field_color(ev["field"])
            alpha = 102 if ev["status"] == "FAIL" else 52
            outline_alpha = 245 if ev["status"] == "FAIL" else 205
            bb = tuple(map(int,b))
            if selected_field == ev["field"]:
                draw.rounded_rectangle(bb, radius=5, fill=_rgba(color,80), outline=_rgba(BLUE,255), width=4)
            else:
                draw.rounded_rectangle(bb, radius=4, fill=_rgba(color,alpha), outline=_rgba(color,outline_alpha), width=2)
            entries.append((b, str(number_map.get(ev["field"],"")), ev["field"], color))
            occupied.append(bb)

    # Small numbered markers, never large text pasted over artwork.
    try:
        font = ImageFont.truetype("DejaVuSans-Bold.ttf", max(13, int(min(base.size)*.018)))
    except Exception:
        font = ImageFont.load_default()
    used = []
    for b, label, name, color in entries:
        x1,y1,x2,y2 = b
        radius = max(12, int(min(base.size)*.010))
        choices = [(x1-radius-5,y1-radius-5),(x2+radius+5,y1-radius-5),(x1-radius-5,y2+radius+5),(x2+radius+5,y2+radius+5)]
        chosen = choices[0]
        for cx,cy in choices:
            cx=max(radius+2,min(base.width-radius-2,cx)); cy=max(radius+2,min(base.height-radius-2,cy))
            rr=(cx-radius,cy-radius,cx+radius,cy+radius)
            if not any(_iou(rr,u)>.02 for u in used):
                chosen=(cx,cy); break
        cx,cy=chosen; rr=(cx-radius,cy-radius,cx+radius,cy+radius); used.append(rr)
        draw.ellipse(rr, fill=(255,255,255,245), outline=_rgba(color,255), width=2)
        tb=draw.textbbox((0,0),label,font=font); tw=tb[2]-tb[0]; th=tb[3]-tb[1]
        draw.text((cx-tw/2,cy-th/2-1),label,fill=_rgba(color,255),font=font)

    return Image.alpha_composite(base, overlay).convert("RGB")


def _image_uri(image, quality=90, max_size=1500):
    if image is None:
        return ""
    if max(image.size) > max_size:
        scale = max_size / max(image.size)
        image = image.resize((max(1,int(image.width*scale)),max(1,int(image.height*scale))), Image.Resampling.LANCZOS)
    buf = io.BytesIO(); image.save(buf, format="JPEG", quality=quality, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


def _registered_org_image(reg):
    if not reg:
        return None
    out_img = reg["output_image"]
    org = reg["org_image"].convert("RGBA")
    size=(max(1,int(org.width*reg["scale"])),max(1,int(org.height*reg["scale"])))
    scaled=org.resize(size,Image.Resampling.LANCZOS)
    canvas=Image.new("RGBA",out_img.size,(255,255,255,255))
    canvas.alpha_composite(scaled,(int(round(reg["tx"])),int(round(reg["ty"]))))
    return canvas.convert("RGB")


def _viewer_html(raw_org_img, registered_org_img, output_img, annotated_img, focus_box=None, height=680):
    raw_org_uri=_image_uri(raw_org_img); reg_org_uri=_image_uri(registered_org_img); out_uri=_image_uri(output_img); ann_uri=_image_uri(annotated_img)
    focus_json = "null" if not focus_box else "[%s,%s,%s,%s]" % tuple(round(float(x),2) for x in focus_box)
    h=int(height)
    return f"""
    <div id='t3viewer' style='height:{h}px;background:#0c1728;border:1px solid #cbd8e6;border-radius:10px;overflow:hidden;position:relative;font-family:Inter,Arial,sans-serif'>
      <div style='height:46px;display:flex;align-items:center;gap:6px;padding:6px 8px;border-bottom:1px solid #22334a;background:#0f2035;color:white;position:absolute;top:0;left:0;right:0;z-index:10'>
        <button class='vbtn active' data-mode='side'>Side by Side</button>
        <button class='vbtn' data-mode='overlay'>Overlay</button>
        <button class='vbtn' data-mode='blink'>Blink</button>
        <span style='width:8px'></span>
        <button class='vbtn' id='minus'>−</button><span id='zoomtxt' style='min-width:40px;text-align:center;font-size:10px'>100%</span><button class='vbtn' id='plus'>+</button>
        <button class='vbtn' id='fit'>FIT</button><button class='vbtn' id='reset'>RESET</button>
        <span id='opacityCtl' style='display:none;align-items:center;gap:5px;font-size:10px;color:#a6b7ca;margin-left:5px'>ORG <input id='opacity' type='range' min='0' max='100' value='50' style='width:95px'><b id='opTxt'>50%</b></span>
        <span style='margin-left:auto;font-size:9px;color:#94a7bb'>Wheel = Zoom • Drag = Pan</span>
      </div>
      <div id='stage' style='position:absolute;inset:46px 0 0 0;overflow:hidden;background:#070f1d;cursor:grab'>
        <div id='canvas' style='position:absolute;left:50%;top:50%;transform-origin:center center;display:flex;align-items:center;gap:16px;will-change:transform'>
          <div id='side' style='display:flex;gap:16px;align-items:center'>
            <div style='background:white'><img src='{raw_org_uri}' style='display:block;max-width:39vw;max-height:585px'></div>
            <div style='background:white'><img src='{ann_uri or out_uri}' style='display:block;max-width:39vw;max-height:585px'></div>
          </div>
          <div id='ov' style='display:none;position:relative;background:white'>
            <img src='{out_uri}' id='ovOut' style='display:block;width:auto;max-width:78vw;max-height:585px'>
            <img src='{reg_org_uri}' id='ovOrg' style='display:block;position:absolute;inset:0;width:100%;height:100%;object-fit:fill;opacity:.5'>
          </div>
          <div id='bl' style='display:none;position:relative;background:white'>
            <img src='{out_uri}' id='blOut' style='display:block;max-width:78vw;max-height:585px'>
            <img src='{reg_org_uri}' id='blOrg' style='display:none;max-width:78vw;max-height:585px'>
          </div>
        </div>
      </div>
    </div>
    <style>
      .vbtn{{background:#10243a;color:#d8e6f3;border:1px solid #2e4c69;border-radius:6px;padding:5px 8px;font-size:10px;cursor:pointer}}
      .vbtn.active{{background:#0b6bb1;border-color:#4aa4e8;color:white}}
      .vbtn:hover{{background:#163552}}
    </style>
    <script>
      const stage=document.getElementById('stage'), canvas=document.getElementById('canvas');
      const side=document.getElementById('side'), ov=document.getElementById('ov'), bl=document.getElementById('bl');
      const op=document.getElementById('opacityCtl'), ovOrg=document.getElementById('ovOrg');
      const blOrg=document.getElementById('blOrg'), blOut=document.getElementById('blOut');
      const zoomTxt=document.getElementById('zoomtxt');
      let mode='side', zoom=1, px=0, py=0, dragging=false, sx=0, sy=0, ox=0, oy=0, timer=null, blink=false, ms=700;
      const focus={focus_json};
      function apply(){{canvas.style.transform='translate(-50%,-50%) translate('+px+'px,'+py+'px) scale('+zoom+')';zoomTxt.textContent=Math.round(zoom*100)+'%';}}
      function hide(){{side.style.display='none';ov.style.display='none';bl.style.display='none';op.style.display='none';}}
      function stop(){{if(timer){{clearInterval(timer);timer=null;}}}}
      function setMode(m){{stop();mode=m;document.querySelectorAll('[data-mode]').forEach(b=>b.classList.toggle('active',b.dataset.mode===m));hide();if(m==='side')side.style.display='flex';if(m==='overlay'){{ov.style.display='block';op.style.display='flex';}}if(m==='blink'){{bl.style.display='block';startBlink();}}apply();}}
      function startBlink(){{blink=false;blOrg.style.display='none';blOut.style.display='block';timer=setInterval(()=>{{blink=!blink;blOrg.style.display=blink?'block':'none';blOut.style.display=blink?'none':'block';}},ms);}}
      document.querySelectorAll('[data-mode]').forEach(b=>b.onclick=()=>setMode(b.dataset.mode));
      document.getElementById('plus').onclick=()=>{{zoom=Math.min(3,zoom+.1);apply();}};
      document.getElementById('minus').onclick=()=>{{zoom=Math.max(.35,zoom-.1);apply();}};
      document.getElementById('fit').onclick=()=>{{zoom=1;px=0;py=0;apply();}};
      document.getElementById('reset').onclick=()=>{{zoom=1;px=0;py=0;apply();}};
      document.getElementById('opacity').oninput=e=>{{ovOrg.style.opacity=(e.target.value/100).toFixed(2);document.getElementById('opTxt').textContent=e.target.value+'%';}};
      stage.addEventListener('wheel',e=>{{e.preventDefault();zoom=Math.max(.35,Math.min(3,zoom+(e.deltaY<0?.1:-.1)));apply();}},{{passive:false}});
      stage.onmousedown=e=>{{dragging=true;sx=e.clientX;sy=e.clientY;ox=px;oy=py;stage.style.cursor='grabbing';}};
      window.addEventListener('mouseup',()=>{{dragging=false;stage.style.cursor='grab';}});
      window.addEventListener('mousemove',e=>{{if(!dragging)return;px=ox+e.clientX-sx;py=oy+e.clientY-sy;apply();}});
      if(focus && Array.isArray(focus)){{
        const x=(focus[0]+focus[2])/2, y=(focus[1]+focus[3])/2;
        const ow={int(output_img.width if output_img else 1)}, oh={int(output_img.height if output_img else 1)};
        const nx=(x/Math.max(1,ow)-.5), ny=(y/Math.max(1,oh)-.5);
        px=-nx*280; py=-ny*330; zoom=1.25;
      }}
      setMode('side');apply();
    </script>
    """


# ============================================================================
# Field browser / UX
# ============================================================================


def _field_search_score(field, query):
    if not query.strip():
        return 0
    q = _norm(query)
    compact = _compact(field)
    tokens = _tokens(field)
    family = _norm(get_field_type(field))
    region = _norm(get_field_region(field))
    score = 0
    if q in _norm(field): score += 100
    if q.replace(" ", "") in compact: score += 80
    if q in family: score += 70
    if q in region: score += 50
    if any(q in t for t in tokens): score += 35
    if any(t.startswith(q[:2]) for t in tokens if q): score += 10
    synonyms = {
        "fiber": {"content": 60, "composition": 60, "material": 55, "fib": 80},
        "fabric": {"content": 60, "composition": 60, "material": 55},
        "care": {"care": 80, "washing": 75, "wash": 75, "wc": 90},
        "country": {"coo": 85, "origin": 75, "made": 70, "min": 75},
        "size": {"size": 90, "osz": 95, "alpha": 50},
        "rn": {"rn": 95, "ca": 75},
        "barcode": {"barcode": 90, "ean": 85, "upc": 85, "gtin": 85},
    }
    for key, families in synonyms.items():
        if key in q:
            for token, boost in families.items():
                if token in compact or token in family:
                    score += boost
    return score


def _filter_fields(fields, query):
    if not query.strip():
        return list(fields)
    ranked = [(field, _field_search_score(field, query)) for field in fields]
    return [field for field, score in sorted(ranked, key=lambda x:(x[1], -fields.index(x[0])), reverse=True) if score > 0]


def _render_field_browser(all_fields, selected_fields, evidence, key_prefix):
    selected_set = set(selected_fields)
    search = st.text_input("Search fields, family or value", placeholder="e.g. care, fiber, country, size, RN...", key=f"{key_prefix}_search")
    filtered = _filter_fields(all_fields, search)
    evidence_by_field = {e["field"]:e for e in evidence}

    left, right = st.columns([1.7, .9], gap="small")
    with left:
        st.markdown(f"<div class='t3-card'><div class='t3-card-title'>Field Comparison ({len(selected_fields)} selected)</div><div class='t3-card-sub'>Search is semantic. Click the small arrow to inspect the complete detected list.</div></div>", unsafe_allow_html=True)
    with right:
        st.markdown(f"<div class='t3-mini-stat'><b>{len(all_fields)}</b><span>Available populated fields</span></div>", unsafe_allow_html=True)

    with st.expander(f"▾  View / Edit Complete Field List  •  {len(selected_fields)} selected", expanded=True):
        if not filtered:
            st.info("No matching fields found.")
        else:
            # Multiselect is the editable source of truth; the rows beneath it are
            # the fully expanded visual list so nothing is hidden behind 2–3 chips.
            current = st.multiselect("Selected fields", filtered if search else all_fields, default=[f for f in selected_fields if f in (filtered if search else all_fields)], key=f"{key_prefix}_multiselect")
            if search and current != selected_fields:
                # Preserve previously selected fields that are currently hidden by search.
                preserved=[f for f in selected_fields if f not in filtered]
                selected_fields=list(dict.fromkeys(preserved+current))
            else:
                selected_fields=current
            st.markdown("<div style='height:4px'></div>", unsafe_allow_html=True)
            for field in filtered:
                ev=evidence_by_field.get(field)
                status=ev["status"] if ev else "AUTO" if field in selected_fields else "INFO"
                locked=bool(ev and ev.get("locked"))
                detected="AUTO" if field in selected_fields and not search else ""
                st.markdown(
                    f"<div class='t3-field-row'><div style='display:flex;justify-content:space-between;align-items:center;gap:8px'><div><div class='t3-field-name'>✓ {html.escape(field)}</div><div class='t3-field-sub'>{html.escape(get_field_type(field))}{(' • '+html.escape(str(get_field_region(field)))) if get_field_region(field) else ''}</div></div><div style='display:flex;gap:5px;align-items:center'>{_status_badge(status)} {('<span class=\"t3-lock\">🔒</span>' if locked else '')}</div></div></div>",
                    unsafe_allow_html=True,
                )
    return selected_fields


# ============================================================================
# Reports / summary
# ============================================================================


def _overall(evidence, static_matches, unaccounted):
    variable = [e for e in evidence if e.get("classification") == "VARIABLE"]
    fails=sum(1 for e in variable if e["status"]=="FAIL")
    reviews=sum(1 for e in variable if e["status"]=="REVIEW")
    overall="FAIL" if fails or unaccounted else "REVIEW" if reviews else "PASS"
    return overall, len(static_matches), len([e for e in variable if e["status"] in {"PASS","FAIL","REVIEW"}]), fails, reviews, len(unaccounted)


def _evidence_dataframe(evidence, static_matches, unaccounted):
    rows=[]
    for e in evidence:
        org=e.get("org_block",{}).get("text", "Not mapped") if e.get("org_block") else "Not mapped"
        rows.append({
            "PDF PAGE":e["page"],"Element / Field":e["field"],"Type":e.get("classification", "VARIABLE"),"ORG Spec":org,
            "Output":e["actual"],"Order Form":e["expected"],"Status":e["status"],
            "Evidence Locked":"YES" if e.get("locked") else "NO","Source":e.get("source", "Tool 1"),"Notes":e["difference"],
        })
    for m in static_matches:
        rows.append({
            "PDF PAGE":m["page"],"Element / Field":"Static Element","Type":"STATIC","ORG Spec":m["org"]["text"],
            "Output":m["output"]["text"],"Order Form":"—","Status":"STATIC","Evidence Locked":"YES","Source":"ORG","Notes":f"Registration confidence {m['score']:.2f}",
        })
    for m in unaccounted:
        rows.append({
            "PDF PAGE":m["page"],"Element / Field":"ORG Element","Type":"MISSING / UNACCOUNTED","ORG Spec":m["expected"],
            "Output":"Not found","Order Form":"—","Status":"MISSING / UNACCOUNTED","Evidence Locked":"YES","Source":"ORG","Notes":m["difference"],
        })
    return pd.DataFrame(rows)


def _build_excel_report(result):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    wb=Workbook(); ws=wb.active; ws.title="Summary"
    dark=PatternFill("solid",fgColor="17324D"); white=Font(color="FFFFFF",bold=True); thin=Side(style="thin",color="D9E4F0")
    ws["A1"]="ORG SPEC + ORDER FORM + OUTPUT QC"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF"); ws["A1"].fill=dark
    summary=result["summary"]
    ws["A2"]="Product Type"; ws["B2"]=result["product_type"]
    labels=["Overall Result","Static Elements","Variable Elements","Issues","Manual Review","Unaccounted ORG Elements","Tool 1 Engine"]
    values=list(summary)+[TOOL1_ENGINE_VERSION]
    for i,(lab,val) in enumerate(zip(labels,values),start=3): ws.cell(i,1,lab); ws.cell(i,2,val)
    comp=_evidence_dataframe(result["evidence"],result["static_matches"],result["unaccounted"])
    detail=wb.create_sheet("Field Comparison")
    if not comp.empty:
        for c,name in enumerate(comp.columns,1): detail.cell(1,c,name).fill=dark; detail.cell(1,c).font=white
        for r,vals in enumerate(comp.itertuples(index=False),2):
            for c,v in enumerate(vals,1): detail.cell(r,c,v); detail.cell(r,c).alignment=Alignment(vertical="top",wrap_text=True)
        status_col=list(comp.columns).index("Status")+1
        fills={"PASS":"16A34A","FAIL":"DC2626","REVIEW":"EA580C","STATIC":"CA8A04","MISSING / UNACCOUNTED":"9333EA"}
        for r in range(2,detail.max_row+1):
            stt=str(detail.cell(r,status_col).value or "")
            if stt in fills: detail.cell(r,status_col).fill=PatternFill("solid",fgColor=fills[stt]); detail.cell(r,status_col).font=white
    reg=wb.create_sheet("Registration")
    headers=["Page","Scale","Scale X","Scale Y","X Offset","Y Offset","Aspect Delta","Warning"]
    for c,h in enumerate(headers,1): reg.cell(1,c,h).fill=dark; reg.cell(1,c).font=white
    for r,(page,rr) in enumerate(result["registrations"],2):
        vals=[page,round(rr["scale"],4),round(rr["scale_x"],4),round(rr["scale_y"],4),round(rr["tx"],1),round(rr["ty"],1),round(rr["aspect_delta"],4),rr["warning"]]
        for c,v in enumerate(vals,1): reg.cell(r,c,v)
    vis=wb.create_sheet("Artwork Visual Validation"); vis["A1"]="Artwork Visual Validation"; vis["A1"].font=Font(size=16,bold=True)
    vis["A3"]="Yellow = Static | Field colors = Variable/Pass | Red = Fail | Orange = Review | Evidence is locked"
    row=5
    for page_no,img in result["annotated_images"].items():
        vis.cell(row,1,f"Page {page_no}"); row+=1
        if img:
            buf=io.BytesIO(); img.save(buf,format="PNG"); buf.seek(0); pic=XLImage(buf); pic.width=min(560,img.width); pic.height=int(img.height*(pic.width/img.width)); vis.add_image(pic,f"A{row}"); row+=max(30,int(pic.height/14)+2)
    for sh in wb.worksheets:
        for col in sh.columns:
            letter=col[0].column_letter
            sh.column_dimensions[letter].width=24 if sh==vis and letter=="A" else min(65,max(12,max(len(str(c.value or "")) for c in col)+2))
        for rowcells in sh.iter_rows():
            for cell in rowcells: cell.alignment=Alignment(vertical="top",wrap_text=True)
    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()


# ============================================================================
# ============================================================================
# WIZARD UI HELPERS
# ============================================================================


def _safe_page_numbers(pages):
    return [int(p.get("page", i + 1)) for i, p in enumerate(pages or [])]


def _status_counts(evidence, static_matches, unaccounted):
    variable = [e for e in evidence if e.get("classification") == "VARIABLE"]
    return {
        "static": len(static_matches),
        "variable": len(variable),
        "pass": sum(1 for e in variable if e.get("status") == "PASS"),
        "fail": sum(1 for e in variable if e.get("status") == "FAIL"),
        "review": sum(1 for e in variable if e.get("status") == "REVIEW"),
        "unaccounted": len(unaccounted),
    }


def _run_tool3_pipeline(result, selected_fields=None):
    """Re-run Tool 1 + Tool 3 enrichment without duplicating Tool 1 rules."""
    selected = list(selected_fields if selected_fields is not None else result["selected_fields"])
    df = result["df"]
    output_pages = result["output_pages"]
    org_pages = result["org_pages"]
    product = result["product_type"]
    mapping = result["mapping"]

    report = _tool1_variable_results(df, output_pages, selected, product, mapping)
    evidence = _build_variable_evidence(df, output_pages, org_pages, report, product)
    static_matches, unaccounted, registrations = _classify_static_and_unaccounted(
        org_pages, output_pages, evidence
    )
    summary = _overall(evidence, static_matches, unaccounted)
    annotated_images = {
        int(p.get("page", i + 1)): _draw_visual_evidence(
            p, evidence, static_matches
        )
        for i, p in enumerate(output_pages)
    }

    result.update({
        "selected_fields": selected,
        "report": report,
        "evidence": evidence,
        "static_matches": static_matches,
        "unaccounted": unaccounted,
        "registrations": registrations,
        "summary": summary,
        "annotated_images": annotated_images,
    })
    return result


def _prepare_run(order_file, org_file, output_file, product_type):
    df = pd.read_csv(order_file) if str(order_file.name).lower().endswith(".csv") else load_excel(order_file)
    if df.empty:
        raise ValueError("The Order Form contains no usable data rows.")

    org_pages = extract_output_pages(org_file)
    output_pages = extract_output_pages(output_file)
    if not org_pages:
        raise ValueError("The ORG Spec could not be read.")
    if not output_pages:
        raise ValueError("The Output artwork could not be read.")

    mapping = _page_row_mapping(df, output_pages)
    detected = _tool1_auto_detect(df, output_pages, product_type, mapping)
    all_fields = [f for f in get_available_fields(df) if not is_admin_field(f)]

    result = {
        "df": df,
        "org_pages": org_pages,
        "output_pages": output_pages,
        "mapping": mapping,
        "all_fields": all_fields,
        "detected_fields": list(detected),
        "selected_fields": list(detected),
        "confirmed_fields": set(),
        "product_type": product_type,
        "files": {
            "order": getattr(order_file, "name", "Order Form"),
            "org": getattr(org_file, "name", "ORG Spec"),
            "output": getattr(output_file, "name", "Output"),
        },
    }
    return _run_tool3_pipeline(result, detected)


def _render_field_detection_table(result, stage, show_full=True):
    """Render a full scrollable field list without relying on Streamlit chips."""
    fields = list(result.get("selected_fields", []))
    evidence = {e["field"]: e for e in result.get("evidence", [])}
    query = st.text_input(
        "Search fields, family or value",
        placeholder="Try: care, fiber, country, size, RN, barcode...",
        key=f"t3_field_search_{stage}",
    )
    filtered = _filter_fields(fields, query)

    if not filtered:
        st.info("No selected fields match the current search.")
        return

    rows = []
    for field in filtered:
        ev = evidence.get(field)
        classification = ev.get("classification", "VARIABLE") if ev else "VARIABLE"
        status = ev.get("status", "AUTO") if ev else "AUTO"
        if classification == "STATIC":
            type_label = "STATIC"
        else:
            type_label = "VARIABLE"
        rows.append({
            "Field": field,
            "Family": get_field_type(field),
            "Region": get_field_region(field) or "—",
            "Type": type_label,
            "Status": status,
            "Source": "AUTO" if field in result.get("detected_fields", []) else "MANUAL",
            "Lock": "🔒" if ev and ev.get("locked") else "—",
        })

    st.dataframe(
        pd.DataFrame(rows),
        use_container_width=True,
        hide_index=True,
        height=430 if show_full else 280,
        column_config={
            "Field": st.column_config.TextColumn("Field", width="medium"),
            "Family": st.column_config.TextColumn("Family", width="small"),
            "Region": st.column_config.TextColumn("Region", width="small"),
            "Type": st.column_config.TextColumn("Type", width="small"),
            "Status": st.column_config.TextColumn("Status", width="small"),
            "Source": st.column_config.TextColumn("Source", width="small"),
            "Lock": st.column_config.TextColumn("Lock", width="small"),
        },
        key=f"t3_field_table_{stage}",
    )


def _render_stage_header(number, title, subtitle):
    st.markdown(
        f"<div class='t3-card' style='margin-bottom:10px'><div style='display:flex;align-items:center;gap:10px'>"
        f"<div style='width:34px;height:34px;border-radius:50%;background:#2563eb;color:#fff;display:flex;align-items:center;justify-content:center;font-weight:900'>{number}</div>"
        f"<div><div class='t3-card-title' style='font-size:17px'>{html.escape(title)}</div>"
        f"<div class='t3-card-sub'>{html.escape(subtitle)}</div></div></div></div>",
        unsafe_allow_html=True,
    )


def _render_visual_stage_preview(result, selected_field, selected_page, stage_key):
    """Show Tool-1-style visual tagging for a selected field during Detect/Review."""
    output_pages = result["output_pages"]
    org_pages = result["org_pages"]
    evidence = result.get("evidence", [])
    static_matches = result.get("static_matches", [])
    page_numbers = _safe_page_numbers(output_pages)
    if not page_numbers:
        return
    if selected_page not in page_numbers:
        selected_page = page_numbers[0]
    idx = page_numbers.index(selected_page)
    out_page = output_pages[idx]
    org_page = org_pages[idx] if idx < len(org_pages) else None
    annotated = _draw_visual_evidence(out_page, evidence, static_matches, selected_field=selected_field)
    reg = _register(org_page, out_page) if org_page else None
    raw_org = reg["org_image"] if reg else (Image.open(io.BytesIO(org_page["image_bytes"])).convert("RGB") if org_page and org_page.get("image_bytes") else None)
    registered_org = _registered_org_image(reg) if reg else raw_org
    out_img = Image.open(io.BytesIO(out_page["image_bytes"])).convert("RGB")

    selected_ev = next((e for e in evidence if e["page"] == selected_page and e["field"] == selected_field), None)
    focus = None
    if selected_ev and selected_ev.get("boxes"):
        boxes = selected_ev["boxes"]
        focus = (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))

    st.components.v1.html(
        _viewer_html(raw_org, registered_org, out_img, annotated or out_img, focus_box=focus, height=520),
        height=540,
        scrolling=False,
    )

    if selected_ev:
        st.markdown(
            f"<div class='t3-role'><b>{html.escape(selected_field)}</b> • "
            f"{_status_badge(selected_ev.get('status','INFO'))} &nbsp; "
            f"{_status_badge(selected_ev.get('classification','VARIABLE'))} &nbsp; "
            f"{'🔒 Evidence locked' if selected_ev.get('locked') else 'Evidence needs review'}<br>"
            f"<b>Order Form:</b> {html.escape(str(selected_ev.get('expected','')))}<br>"
            f"<b>Output:</b> {html.escape(str(selected_ev.get('actual','')))}<br>"
            f"<b>Finding:</b> {html.escape(str(selected_ev.get('difference','—')))}</div>",
            unsafe_allow_html=True,
        )


def _render_navigation(back_step=None, next_label="Next", next_key=None, next_disabled=False):
    left, right = st.columns([1, 1])
    with left:
        if back_step is not None:
            if st.button("← Back", width="stretch", key=f"{next_key}_back"):
                st.session_state["t3_stage"] = back_step
                st.rerun()
    with right:
        if st.button(f"{next_label} →", type="primary", width="stretch", key=next_key, disabled=next_disabled):
            return True
    return False


# ============================================================================
# Main — five-stage wizard
# ============================================================================


def main():
    _tool3_css()

    if "t3_reset_id" not in st.session_state:
        st.session_state["t3_reset_id"] = 0
    if "t3_stage" not in st.session_state:
        st.session_state["t3_stage"] = 1
    if "t3_result" not in st.session_state:
        st.session_state["t3_result"] = None
    if "t3_focus_field" not in st.session_state:
        st.session_state["t3_focus_field"] = None
    if "t3_confirmed_fields" not in st.session_state:
        st.session_state["t3_confirmed_fields"] = set()

    if st.session_state.get("t3_version") != TOOL3_VERSION:
        st.session_state["t3_stage"] = 1
        st.session_state["t3_result"] = None
        st.session_state["t3_focus_field"] = None
        st.session_state["t3_confirmed_fields"] = set()
        st.session_state["t3_version"] = TOOL3_VERSION

    reset = st.session_state["t3_reset_id"]
    stage = int(st.session_state["t3_stage"])

    _render_top()
    _render_stepper(stage)

    # ------------------------------------------------------------------
    # Stage-specific input surface
    # ------------------------------------------------------------------
    if stage == 1:
        st.markdown("<div class='t3-section-title'>Project Inputs</div>", unsafe_allow_html=True)
        c1, c2, c3, c4 = st.columns([1.05, 1.05, 1.05, .95], gap="small")

        with c1:
            st.markdown("<div class='t3-card'><div class='t3-upload-icon'>📊</div><div class='t3-card-title'>Order Form</div><div class='t3-card-sub'>Source of variable data</div>", unsafe_allow_html=True)
            order_file = st.file_uploader("Order Form", type=["xlsx", "xls", "csv"], label_visibility="collapsed", key=f"t3_order_{reset}")
            if order_file:
                st.markdown(f"<div class='t3-fileline'>{html.escape(order_file.name)} ✓</div>", unsafe_allow_html=True)
            st.markdown("<div class='t3-role'><b>Used for:</b> variable fields such as FIB, WC, RN, COO, Size, Brand, Barcode, etc.</div></div>", unsafe_allow_html=True)

        with c2:
            st.markdown("<div class='t3-card'><div class='t3-upload-icon'>📐</div><div class='t3-card-title'>ORG Spec</div><div class='t3-card-sub'>Approved content + presentation baseline</div>", unsafe_allow_html=True)
            org_file = st.file_uploader("ORG Spec", type=["pdf", "jpg", "jpeg", "png"], label_visibility="collapsed", key=f"t3_org_{reset}")
            if org_file:
                st.markdown(f"<div class='t3-fileline'>{html.escape(org_file.name)} ✓</div>", unsafe_allow_html=True)
            st.markdown("<div class='t3-role'><b>Used for:</b> wording, case, punctuation, position, layout and static reference.</div></div>", unsafe_allow_html=True)

        with c3:
            st.markdown("<div class='t3-card'><div class='t3-upload-icon'>🖼️</div><div class='t3-card-title'>Output</div><div class='t3-card-sub'>Final artwork to validate</div>", unsafe_allow_html=True)
            output_file = st.file_uploader("Output", type=["pdf", "jpg", "jpeg", "png"], label_visibility="collapsed", key=f"t3_output_{reset}")
            if output_file:
                st.markdown(f"<div class='t3-fileline'>{html.escape(output_file.name)} ✓</div>", unsafe_allow_html=True)
            st.markdown("<div class='t3-role'><b>Checked against:</b> ORG baseline + Tool 1 Order Form validation.</div></div>", unsafe_allow_html=True)

        with c4:
            st.markdown("<div class='t3-card'><div class='t3-card-title'>Product Type</div><div class='t3-card-sub'>Controls variable-data presentation behavior</div>", unsafe_allow_html=True)
            product = st.radio("Product Type", ["PFL", "HTL", "Other"], horizontal=True, key=f"t3_product_{reset}")
            st.markdown("<div class='t3-role'><b>PFL:</b> Order Form presentation. <b>HTL / Other:</b> ORG presentation baseline.</div></div>", unsafe_allow_html=True)

        ready = bool(order_file and org_file and output_file)
        if not ready:
            st.info("Upload Order Form, ORG Spec and Output to start the five-stage QC workflow.")
    else:
        result_preview = st.session_state.get("t3_result")
        if not result_preview:
            st.session_state["t3_stage"] = 1
            st.rerun()
            return
        product = result_preview["product_type"]
        order_name = result_preview["files"].get("order", "Order Form")
        org_name = result_preview["files"].get("org", "ORG Spec")
        output_name = result_preview["files"].get("output", "Output")
        st.markdown(
            f"<div class='t3-card' style='margin-bottom:10px'><div style='display:flex;gap:7px;flex-wrap:wrap;align-items:center'>"
            f"<span class='t3-badge' style='background:#dcfce7;color:#15803d'>✓ Order Form</span><span style='font-size:9px;color:#64748b'>{html.escape(order_name)}</span>"
            f"<span style='color:#cbd5e1'>•</span><span class='t3-badge' style='background:#e0e7ff;color:#4338ca'>✓ ORG Spec</span><span style='font-size:9px;color:#64748b'>{html.escape(org_name)}</span>"
            f"<span style='color:#cbd5e1'>•</span><span class='t3-badge' style='background:#dbeafe;color:#1e40af'>✓ Output</span><span style='font-size:9px;color:#64748b'>{html.escape(output_name)}</span>"
            f"<span style='margin-left:auto' class='t3-badge' style='background:#eff6ff;color:#1d4ed8'>Product: {html.escape(product)}</span>"
            f"</div></div>",
            unsafe_allow_html=True,
        )
        order_file = org_file = output_file = None
        ready = True

    # ------------------------------------------------------------------
    # STAGE 1 — Upload / initialize
    # ------------------------------------------------------------------
    if stage == 1:
        _render_stage_header(1, "Upload Files & Select Product Type", "Provide the three source files and choose the product type before analysis starts.")
        st.markdown(
            "<div class='t3-hero-result'><div class='t3-hero-title'>What happens next?</div>"
            "<div class='t3-hero-sub'>Tool 1 will detect relevant Order Form fields. Tool 3 will then analyze ORG structure, compare Output, lock evidence, and prepare the visual QC workspace.</div></div>",
            unsafe_allow_html=True,
        )
        if _render_navigation(next_label="Start Auto Detect", next_key=f"t3_stage1_next_{reset}"):
            try:
                with st.spinner("Preparing files and running Tool 1 Auto Detect…"):
                    result = _prepare_run(order_file, org_file, output_file, product)
                    st.session_state["t3_result"] = result
                    st.session_state["t3_confirmed_fields"] = set()
                    st.session_state["t3_stage"] = 2
                    st.session_state["t3_focus_field"] = result["selected_fields"][0] if result["selected_fields"] else None
                st.rerun()
            except Exception as exc:
                st.error(f"QC setup could not be completed: {exc}")
        return

    result = st.session_state.get("t3_result")
    if not result:
        st.session_state["t3_stage"] = 1
        st.rerun()
        return

    # ------------------------------------------------------------------
    # STAGE 2 — Auto Detect + visual tagging
    # ------------------------------------------------------------------
    if stage == 2:
        _render_stage_header(2, "Auto Detect Fields", "Review what Tool 1 detected and see the exact artwork regions before confirming them.")
        counts = _status_counts(result["evidence"], result["static_matches"], result["unaccounted"])
        stats = st.columns(4)
        with stats[0]: st.markdown(f"<div class='t3-mini-stat'><b>{len(result['detected_fields'])}</b><span>Fields detected</span></div>", unsafe_allow_html=True)
        with stats[1]: st.markdown(f"<div class='t3-mini-stat'><b>{sum(1 for e in result['evidence'] if e.get('locked'))}</b><span>Evidence mapped</span></div>", unsafe_allow_html=True)
        with stats[2]: st.markdown(f"<div class='t3-mini-stat'><b>{counts['fail']}</b><span>Current fails</span></div>", unsafe_allow_html=True)
        with stats[3]: st.markdown(f"<div class='t3-mini-stat'><b>{counts['review']}</b><span>Review items</span></div>", unsafe_allow_html=True)

        left, right = st.columns([1.15, 1.85], gap="small")
        with left:
            _render_field_detection_table(result, 2, show_full=True)
            selected_options = result["selected_fields"]
            if selected_options:
                selected_field = st.selectbox("Inspect detected field", selected_options, index=max(0, selected_options.index(st.session_state.get("t3_focus_field")) if st.session_state.get("t3_focus_field") in selected_options else 0), key=f"t3_detect_field_{reset}")
                st.session_state["t3_focus_field"] = selected_field
            else:
                selected_field = None
                st.warning("Auto Detect did not find a usable populated field. You can add fields manually in Step 3.")

        with right:
            if selected_field:
                page_numbers = _safe_page_numbers(result["output_pages"])
                page = st.selectbox("Artwork Page", page_numbers, key=f"t3_detect_page_{reset}")
                _render_visual_stage_preview(result, selected_field, page, "detect")
            else:
                st.info("Select a field to see its visual tag on the artwork.")

        st.markdown(
            "<div class='t3-role'><b>Visual review:</b> the highlight shows the evidence used by the comparison engine. Green/field-color = usable match, red = mismatch, amber = presentation/review issue, yellow = exact ORG static evidence.</div>",
            unsafe_allow_html=True,
        )

        back, next_col = st.columns([1,1])
        with back:
            if st.button("← Back to Upload", width="stretch", key=f"t3_stage2_back_{reset}"):
                st.session_state["t3_stage"] = 1
                st.rerun()
        with next_col:
            if st.button("Review & Edit Fields →", type="primary", width="stretch", key=f"t3_stage2_next_{reset}"):
                st.session_state["t3_stage"] = 3
                st.rerun()
        return

    # ------------------------------------------------------------------
    # STAGE 3 — Review / edit / confirm detected fields
    # ------------------------------------------------------------------
    if stage == 3:
        _render_stage_header(3, "Review & Edit Detected Fields", "Confirm, remove or add fields. Selected evidence remains locked so another field cannot reuse it.")
        left, right = st.columns([1.3, 1.7], gap="small")

        with left:
            st.markdown(f"<div class='t3-card'><div class='t3-card-title'>Detected Fields ({len(result['selected_fields'])})</div><div class='t3-card-sub'>The full list is scrollable. Search by field name, family or semantic meaning.</div></div>", unsafe_allow_html=True)
            current = st.multiselect(
                "Fields participating in QC",
                result["all_fields"],
                default=[f for f in result["selected_fields"] if f in result["all_fields"]],
                key=f"t3_review_multiselect_{reset}",
            )
            st.caption(f"Selected: {len(current)} fields • Auto detected: {len(result['detected_fields'])}")
            _render_field_detection_table({**result, "selected_fields": current}, 3, show_full=True)
            if current != result["selected_fields"]:
                st.info("Field selection has changed. The comparison will be recalculated when you continue.")

            page_numbers = _safe_page_numbers(result["output_pages"])
            mapping_changed = False
            with st.expander("▾ Page Mapping", expanded=len(page_numbers) > 1):
                opts = list(range(len(result["df"])))
                labels = [f"Excel Row {i + 2}" for i in opts]
                for i, pn in enumerate(page_numbers):
                    old = int(result["mapping"].get(pn, min(i, len(result["df"]) - 1)))
                    chosen = st.selectbox(
                        f"PDF Page {pn}", opts,
                        index=max(0, min(old, len(opts) - 1)),
                        format_func=lambda x, labels=labels: labels[x],
                        key=f"t3_review_map_{reset}_{pn}",
                    )
                    if chosen != old:
                        mapping_changed = True
                    result["mapping"][pn] = chosen

        with right:
            inspect_fields = current or result["selected_fields"]
            if inspect_fields:
                focus_default = st.session_state.get("t3_focus_field")
                idx_default = inspect_fields.index(focus_default) if focus_default in inspect_fields else 0
                selected_field = st.selectbox("Inspect selected field", inspect_fields, index=idx_default, key=f"t3_review_field_{reset}")
                st.session_state["t3_focus_field"] = selected_field
                page_numbers = _safe_page_numbers(result["output_pages"])
                selected_page = st.selectbox("Artwork Page", page_numbers, key=f"t3_review_page_{reset}")
                _render_visual_stage_preview(result, selected_field, selected_page, "review")
                ev = next((e for e in result["evidence"] if e["field"] == selected_field and e["page"] == selected_page), None)
                confirm_key = f"t3_confirm_{reset}_{selected_page}_{selected_field}"
                confirmed = st.checkbox("I confirm this detected field/evidence is correct", value=(selected_field in st.session_state["t3_confirmed_fields"]), key=confirm_key)
                if confirmed:
                    st.session_state["t3_confirmed_fields"].add(selected_field)
                else:
                    st.session_state["t3_confirmed_fields"].discard(selected_field)
                if ev:
                    st.markdown(f"<div class='t3-role'><b>Evidence state:</b> {'🔒 Confirmed + locked' if selected_field in st.session_state['t3_confirmed_fields'] else 'Review pending'}<br><b>Classification:</b> {html.escape(ev.get('classification','VARIABLE'))}<br><b>Result:</b> {_status_badge(ev.get('status','INFO'))}</div>", unsafe_allow_html=True)
            else:
                st.warning("No fields selected. Add or select at least one field before continuing.")

        back, next_col = st.columns([1,1])
        with back:
            if st.button("← Back to Auto Detect", width="stretch", key=f"t3_stage3_back_{reset}"):
                st.session_state["t3_stage"] = 2
                st.rerun()
        with next_col:
            if st.button("Proceed to Comparison →", type="primary", width="stretch", disabled=not bool(current), key=f"t3_stage3_next_{reset}"):
                with st.spinner("Re-running Tool 1 validation for the confirmed field set…"):
                    result["selected_fields"] = list(current)
                    _run_tool3_pipeline(result, current)
                    result["confirmed_fields"] = set(st.session_state.get("t3_confirmed_fields", set()))
                    st.session_state["t3_result"] = result
                    st.session_state["t3_stage"] = 4
                    st.rerun()
        return

    # ------------------------------------------------------------------
    # STAGE 4 — comparison workspace
    # ------------------------------------------------------------------
    if stage == 4:
        _render_stage_header(4, "Comparison View", "Inspect the registered ORG ↔ Output artwork, click findings, and verify field evidence.")
        overall, static_count, var_count, fail_count, review_count, unaccounted_count = result["summary"]
        page_numbers = _safe_page_numbers(result["output_pages"])

        controls = st.columns([.8, 1.1, 1.2, .9])
        with controls[0]:
            selected_page = st.selectbox("Page", page_numbers, key=f"t3_cmp_page_{reset}")
        with controls[1]:
            filter_choice = st.selectbox("Findings", ["All", "FAIL", "REVIEW", "PASS", "STATIC"], key=f"t3_cmp_filter_{reset}")
        with controls[2]:
            st.markdown(f"<div class='t3-mini-stat'><b>{overall}</b><span>{static_count} Static • {var_count} Variable • {fail_count} Fail</span></div>", unsafe_allow_html=True)
        with controls[3]:
            if st.button("↻ Recalculate", width="stretch", key=f"t3_recalc_{reset}"):
                with st.spinner("Re-running Tool 1 + Tool 3 classification…"):
                    _run_tool3_pipeline(result, result["selected_fields"])
                    st.rerun()

        idx = page_numbers.index(selected_page)
        out_page = result["output_pages"][idx]
        org_page = result["org_pages"][idx] if idx < len(result["org_pages"]) else None
        reg = _register(org_page, out_page) if org_page else None
        raw_org = reg["org_image"] if reg else None
        registered_org = _registered_org_image(reg) if reg else raw_org
        out_img = Image.open(io.BytesIO(out_page["image_bytes"])).convert("RGB")
        focus_field = st.session_state.get("t3_focus_field")
        annotated = _draw_visual_evidence(out_page, result["evidence"], result["static_matches"], selected_field=focus_field)

        focus_box = None
        if focus_field:
            ev_focus = next((e for e in result["evidence"] if e["field"] == focus_field and e["page"] == selected_page), None)
            if ev_focus and ev_focus.get("boxes"):
                bxs = ev_focus["boxes"]
                focus_box = (min(b[0] for b in bxs), min(b[1] for b in bxs), max(b[2] for b in bxs), max(b[3] for b in bxs))

        st.components.v1.html(_viewer_html(raw_org, registered_org, out_img, annotated or out_img, focus_box=focus_box, height=680), height=700, scrolling=False)
        if reg:
            st.caption(f"Registered ✓ • Scale {reg['scale']:.3f} • X {reg['tx']:.1f} • Y {reg['ty']:.1f}" + (f" • {reg['warning']}" if reg.get('warning') else ""))

        viewer_left, viewer_right = st.columns([2.0, 1.0], gap="small")
        page_evidence = [e for e in result["evidence"] if e["page"] == selected_page]
        if filter_choice == "FAIL":
            display_findings = [e for e in page_evidence if e.get("status") == "FAIL"]
        elif filter_choice == "REVIEW":
            display_findings = [e for e in page_evidence if e.get("status") == "REVIEW"]
        elif filter_choice == "PASS":
            display_findings = [e for e in page_evidence if e.get("status") == "PASS"]
        elif filter_choice == "STATIC":
            display_findings = [e for e in page_evidence if e.get("classification") == "STATIC"]
        else:
            display_findings = page_evidence

        with viewer_right:
            st.markdown("<div class='t3-card-title'>Findings</div>", unsafe_allow_html=True)
            find_search = st.text_input("Search findings", placeholder="field, family or value...", key=f"t3_find_search_{reset}")
            if find_search.strip():
                q = _norm(find_search)
                display_findings = [e for e in display_findings if q in _norm(e["field"]) or q in _norm(e.get("actual", "")) or q in _norm(e.get("expected", "")) or q in _norm(e.get("field_type", ""))]
            if not display_findings:
                st.info("No findings match this filter.")
            for e in display_findings:
                selected = e["field"] == focus_field
                col = st.columns([4.3, 1.0])
                with col[0]:
                    if st.button(f"{e['field']}", width="stretch", key=f"t3_find_{reset}_{selected_page}_{e['field']}"):
                        st.session_state["t3_focus_field"] = e["field"]
                        st.rerun()
                with col[1]:
                    st.markdown(_status_badge(e.get("status", e.get("classification", "INFO"))), unsafe_allow_html=True)
                st.markdown(f"<div class='t3-finding-detail'>{html.escape(str(e.get('actual',''))[:95])} {'🔒' if e.get('locked') else ''}</div>", unsafe_allow_html=True)
            st.markdown("<div class='t3-role'><b>Click a finding</b> to focus and zoom toward its locked evidence. Static findings are shown separately from variable findings.</div>", unsafe_allow_html=True)

        with viewer_left:
            chosen = next((e for e in page_evidence if e["field"] == focus_field), None)
            if chosen:
                st.markdown("<div class='t3-section-title'>Selected Finding</div>", unsafe_allow_html=True)
                c1, c2, c3 = st.columns([1.0, 1.1, 1.2])
                with c1:
                    st.markdown(f"<div class='t3-card'><div class='t3-card-title'>{html.escape(chosen['field'])}</div><div style='margin-top:5px'>{_status_badge(chosen.get('classification','VARIABLE'))} {_status_badge(chosen.get('status','INFO'))}</div><div class='t3-card-sub' style='margin-top:6px'>Family: {html.escape(chosen.get('field_type',''))}<br>Region: {html.escape(str(chosen.get('region') or '—'))}<br>{'🔒 Evidence locked' if chosen.get('locked') else 'Evidence not locked'}</div></div>", unsafe_allow_html=True)
                with c2:
                    st.markdown(f"<div class='t3-card'><div class='t3-card-title'>Values</div><div class='t3-card-sub' style='margin-top:6px'><b>Order Form</b><br>{html.escape(chosen.get('expected',''))}<br><br><b>Output</b><br>{html.escape(chosen.get('actual',''))}</div></div>", unsafe_allow_html=True)
                with c3:
                    st.markdown(f"<div class='t3-card'><div class='t3-card-title'>Why?</div><div class='t3-card-sub' style='margin-top:6px'>{html.escape(chosen.get('difference','—'))}</div></div>", unsafe_allow_html=True)
            else:
                st.info("Select a finding on the right to inspect its evidence and reason.")

        back, next_col = st.columns([1,1])
        with back:
            if st.button("← Back to Review", width="stretch", key=f"t3_stage4_back_{reset}"):
                st.session_state["t3_stage"] = 3
                st.rerun()
        with next_col:
            if st.button("View QC Results →", type="primary", width="stretch", key=f"t3_stage4_next_{reset}"):
                st.session_state["t3_stage"] = 5
                st.rerun()
        return

    # ------------------------------------------------------------------
    # STAGE 5 — final QC result / report
    # ------------------------------------------------------------------
    if stage == 5:
        _render_stage_header(5, "QC Results", "Review the final decision, investigate issues and download the complete QC report.")
        overall, static_count, var_count, fail_count, review_count, unaccounted_count = result["summary"]
        hero_bg, hero_fg = (RED_BG, RED) if overall == "FAIL" else (ORANGE_BG, ORANGE) if overall == "REVIEW" else (GREEN_BG, GREEN)
        st.markdown(
            f"<div class='t3-hero-result' style='background:{hero_bg};border-color:{hero_fg}55'>"
            f"<div class='t3-hero-title' style='font-size:22px;color:{hero_fg}'>{overall}</div>"
            f"<div class='t3-hero-sub'>{static_count} Static • {var_count} Variable • {fail_count} Fail • {review_count} Review • {unaccounted_count} Unaccounted</div></div>",
            unsafe_allow_html=True,
        )

        metrics = st.columns(5)
        metric_values = [(static_count, "Static"), (var_count, "Variable"), (sum(1 for e in result["evidence"] if e.get("classification") == "VARIABLE" and e.get("status") == "PASS"), "Pass"), (fail_count, "Fail"), (review_count, "Review")]
        for col, (value, label) in zip(metrics, metric_values):
            with col:
                st.markdown(f"<div class='t3-mini-stat'><b>{value}</b><span>{label}</span></div>", unsafe_allow_html=True)

        issues = [e for e in result["evidence"] if e.get("classification") == "VARIABLE" and e.get("status") in {"FAIL", "REVIEW"}]
        if issues or result["unaccounted"]:
            st.markdown("<div class='t3-section-title'>Issues Requiring Attention</div>", unsafe_allow_html=True)
            for e in issues:
                st.markdown(
                    f"<div class='t3-field-row'><div style='display:flex;justify-content:space-between;gap:8px'><div class='t3-field-name'>{html.escape(e['field'])}</div>{_status_badge(e['status'])}</div><div class='t3-field-sub'>Page {e['page']} • Expected: {html.escape(e['expected'])} • Found: {html.escape(e['actual'])}<br>{html.escape(e['difference'])}</div></div>",
                    unsafe_allow_html=True,
                )
            for u in result["unaccounted"]:
                st.markdown(
                    f"<div class='t3-field-row'><div style='display:flex;justify-content:space-between;gap:8px'><div class='t3-field-name'>ORG Element</div>{_status_badge('MISSING / UNACCOUNTED')}</div><div class='t3-field-sub'>Page {u['page']} • {html.escape(u['expected'])}<br>{html.escape(u['difference'])}</div></div>",
                    unsafe_allow_html=True,
                )
        else:
            st.success("No unresolved variable failures, presentation failures, reviews or unaccounted ORG elements were found.")

        st.markdown("<div class='t3-section-title'>Download QC Report</div>", unsafe_allow_html=True)
        report_bytes = _build_excel_report(result)
        st.download_button(
            "⬇ Download QC Excel Report",
            data=report_bytes,
            file_name="ORG_OrderForm_Output_QC_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
            key=f"t3_download_{reset}",
        )

        st.markdown(
            f"<div class='t3-role'><b>Engine:</b> Tool 1 {html.escape(TOOL1_ENGINE_VERSION)}<br>"
            f"Tool 3 adds ORG baseline, static/variable classification, presentation checks, evidence locking, registration and visual reporting.</div>",
            unsafe_allow_html=True,
        )

        back, new_col = st.columns([1,1])
        with back:
            if st.button("← Back to Comparison", width="stretch", key=f"t3_stage5_back_{reset}"):
                st.session_state["t3_stage"] = 4
                st.rerun()
        with new_col:
            if st.button("＋ Start New QC", width="stretch", key=f"t3_new_{reset}"):
                st.session_state["t3_reset_id"] += 1
                st.session_state["t3_stage"] = 1
                st.session_state["t3_result"] = None
                st.session_state["t3_focus_field"] = None
                st.session_state["t3_confirmed_fields"] = set()
                st.rerun()


if __name__ == "__main__":
    main()
